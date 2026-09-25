#!/usr/bin/env python3
"""Build Review1_Slides_TrekLink.pptx from the official Review 1 template, and render the outline.

    python3 _handoff/tools/build_deck.py <Review1_Slide_Template.pptx>

Slides 1 to 6 are the template's own slides, filled in place, in the template's order. Each Main
Flow uses the template's slide 7 pattern for its steps, followed by one slide holding its swimlane
figure, upright. The closing slides reuse the template's two-box and table slides. The template's
TIP bar carries the presenter's name on every slide, because every member presents.
"""
import copy
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Mm, Pt

REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "_handoff/outbound/capstone/Documents/reports"
FIGS = dict(l.split() for l in (OUT / "assets/mermaid/figures.txt").read_text().splitlines())
FOOTER = "TrekLink · Capstone Project, Review 1 · GFA26SE55"

KHOA, LONGLP, HOANG, LONGNN, TAN = ("Đỗ Đăng Khoa", "Lâm Phi Long", "Trần Khải Hoàng",
                                   "Nguyễn Ngọc Long", "Nguyễn Bá Tân")

MF = [
    ("mf01", "Main Flow 1: Booking to Rental to Trip Preparation", TAN,
     [("Book", "Customer books a package; one device is held for 10 minutes."),
      ("Confirm", "Operator confirms, seeing device and Guide conflicts first."),
      ("Allocate", "Operator allocates devices, loads the fleet key, generates the agreement."),
      ("Sign and pay", "Customer signs; deposit paid; Guide confirms handover and checklist."),
      ("Check out", "Devices move to Rented; the trip is ready to start.")],
     "Two customers take the last device at once: one is told it is no longer available."),
    ("mf02", "Main Flow 2: Field Data to Offline Gateway to Cloud Sync", KHOA,
     [("Broadcast", "Device sends SOS, position or telemetry over the LoRa mesh."),
      ("Classify", "Event tagged P0 to P3 and queued durably before any send."),
      ("Hold", "Uplink down: held on the device (Stage B) and at basecamp (Stage C)."),
      ("Flush", "Uplink back: strict priority order, SOS first."),
      ("Deduplicate", "Backend keys it by eventId; a replay is a no-op.")],
     "The same packet delivered 10 times creates 0 duplicate incidents."),
    ("mf03", "Main Flow 3: SOS to Incident to Emergency Response", HOANG,
     [("SOS", "Button, gesture or fall detection triggers the SOS."),
      ("One incident", "Backend opens one incident per episode, or appends to it."),
      ("Alert", "Guide and Operators alerted within 2 seconds."),
      ("Acknowledge", "First acknowledgement takes ownership and starts the clock."),
      ("Resolve", "In Progress, Resolved, Closed: each with actor, time and note.")],
     "The SOS text frame is lost: beacon cadence raises a Suspected incident."),
    ("mf04", "Main Flow 4: Real-Time Trip Monitoring", LONGNN,
     [("Ingest", "Positions and telemetry arrive through Main Flow 2."),
      ("Update", "Backend updates position, battery and last-seen."),
      ("Scope", "Server pushes only what each role may see."),
      ("Show", "Map shows status, battery, last-seen and incidents."),
      ("Age", "Silent devices and gateways turn visibly stale.")],
     "A Guide asks for another Guide's trip and is refused on the server."),
    ("mf05", "Main Flow 5: Return to Inspection to Billing to Maintenance", LONGLP,
     [("Collect", "Guide collects every device and ends the trip."),
      ("Check in", "Operator checks devices in and inspects them."),
      ("Charge", "Base, late and damage fees computed; deposit applied first."),
      ("Settle", "Balance paid or refunded in sandbox; rental closes."),
      ("Dispose", "Device back to Available, to Maintenance, or Retired.")],
     "Payment fails: the rental stays open with the balance visible."),
]


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def shape(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    raise KeyError(name)


def set_text(sh, lines, size=None):
    """Replace a text frame's text, keeping the first run's formatting."""
    if isinstance(lines, str):
        lines = [lines]
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0] if p0.runs else None
    rpr = copy.deepcopy(r0._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")) if r0 else None
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        p = p0 if i == 0 else tf.add_paragraph()
        if i and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        run = p.add_run()
        run.text = line
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
        if size:
            run.font.size = Pt(size)


def duplicate(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for s in list(new.shapes):
        s._element.getparent().remove(s._element)
    for el in src.shapes._spTree:
        if el.tag.endswith(("}sp", "}grpSp", "}graphicFrame", "}cxnSp", "}pic")):
            new.shapes._spTree.append(copy.deepcopy(el))
    return new


def move(prs, slide, index):
    ids = prs.slides._sldIdLst
    el = [e for e in ids if e.rId == [r for r, p in prs.part.rels.items()
                                      if getattr(p, "target_part", None) is slide.part][0]][0]
    ids.remove(el)
    ids.insert(index, el)


def delete(prs, slide):
    ids = prs.slides._sldIdLst
    for e in list(ids):
        if prs.part.related_part(e.rId) is slide.part:
            prs.part.drop_rel(e.rId)
            ids.remove(e)


def common(slide, number, title, badge, presenter):
    set_text(shape(slide, "Text 1"), str(number))
    set_text(shape(slide, "Text 2"), title, size=24 if len(title) > 40 else None)
    set_text(shape(slide, "Text 4"), badge)
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.startswith("TIP"):
            set_text(s, f"Presenter: {presenter}")
        if s.has_text_frame and s.text_frame.text.startswith("Capstone Project"):
            set_text(s, FOOTER)


def fill_table(slide, rows, widths=None):
    tbl_shape = [s for s in slide.shapes if s.has_table][0]
    t = tbl_shape.table
    while len(t.rows) < len(rows):
        t._tbl.append(copy.deepcopy(t.rows[len(t.rows) - 1]._tr))
    while len(t.rows) > len(rows):
        t._tbl.remove(t.rows[len(t.rows) - 1]._tr)
    for r, vals in zip(t.rows, rows):
        for c, v in zip(r.cells, vals):
            for p in list(c.text_frame.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            p = c.text_frame.paragraphs[0]
            runs = p.runs
            if runs:
                runs[0].text = v
                for x in runs[1:]:
                    x._r.getparent().remove(x._r)
            else:
                p.add_run().text = v
            for x in p.runs:
                x.font.size = Pt(12)
    if widths:
        for col, w in zip(t.columns, widths):
            col.width = Mm(w)
    return t


# ---------------------------------------------------------------------------
def build(template: Path) -> Path:
    prs = Presentation(str(template))
    s1, s2, s3, s4, s5, s6, s7, s8 = list(prs.slides)

    # 1 Title
    set_text(shape(s1, "Text 5"), "TrekLink Operations Platform", size=36)
    set_text(shape(s1, "Text 8"), [f"{KHOA}, Team Leader", f"{LONGLP}, Member", f"{HOANG}, Member",
                                   f"{LONGNN}, Member", f"{TAN}, Member"], size=12)
    set_text(shape(s1, "Text 11"), ["Mentor: Đặng Ngọc Minh Đức", "FA26SE159 · GFA26SE55"])
    tb = s1.shapes.add_textbox(Mm(14), Mm(96), Mm(284), Mm(10))
    tb.text_frame.text = "Off-grid trekking safety and operations over a LoRa mesh · Presenter: " + KHOA
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    # 2 Context
    common(s2, 2, "Context", "WHY IT MATTERS", KHOA)
    set_text(shape(s2, "Text 6"), "The situation")
    set_text(shape(s2, "Text 7"), [
        "Trekking agencies run multi-day routes through Vietnam's cellular dead zones.",
        "Affected: operators at base, guides in the field, customers on the trail.",
        "Summer 2026: our LoRa mesh firmware sends SOS, fall alerts and positions without coverage.",
        "Nothing yet carries that data to the people who must act on it.",
        "Scope: web system, gateway, targeted firmware work."], size=12)
    set_text(shape(s2, "Text 9"), "A system, with platform seams")
    set_text(shape(s2, "Text 10"), [
        "We build one operations system for one agency.",
        "Beyond that, four seams:",
        "· new field transports plug in (D-007)",
        "· every business rule is configuration (D-015)",
        "· stock Meshtastic devices and apps keep working (D-019)",
        "· staff roles extend as data"], size=12)

    # 3 Problems
    common(s3, 3, "Existing Situation & Problems", "PAIN POINTS", LONGLP)
    probs = [("Field events are lost", "The node's queue holds 16 events in RAM and drops the oldest: in an outage the SOS goes first (MQTT.cpp:821)."),
             ("An SOS goes nowhere", "A device SOS raises a local radio alarm. No one is assigned, nothing is recorded."),
             ("No operational records", "Devices, rentals and trips are tracked in no system; coordination is by phone."),
             ("Impact", "Response depends on who notices; no audit trail after an emergency; custody disputed by memory.")]
    for (t, d), (a, b) in zip(probs, [("Text 8", "Text 9"), ("Text 13", "Text 14"), ("Text 18", "Text 19"), ("Text 23", "Text 24")]):
        set_text(shape(s3, a), t)
        set_text(shape(s3, b), d, size=12)

    # 4 Solution
    common(s4, 4, "Proposed Solution", "OUR APPROACH", KHOA)
    set_text(shape(s4, "Text 6"), "In one sentence")
    set_text(shape(s4, "Text 7"), [
        "A web operations system plus gateway that delivers every field event exactly once, turns each SOS into one owned incident, and runs the fleet, rentals and trips.",
        "Lost events: durable priority queues, SOS first, eventId deduplication.",
        "SOS goes nowhere: one incident per episode, alert in 2 s, full audit trail.",
        "No records: booking, rental, 7-state device lifecycle, sandbox billing."], size=12)
    set_text(shape(s4, "Text 9"), "Approach and boundary")
    set_text(shape(s4, "Text 10"), [
        "React web app · NestJS backend · Node.js gateway · MQTT · PostgreSQL · MapLibre on Goong Maps.",
        "Out of scope: mesh-stack rearchitecture, native apps, real payments, route recommendation, localization, hardware certification, multi-tenancy."], size=12)

    # 5 Features
    common(s5, 5, "Key Features", "PRODUCT SCOPE", HOANG)
    feats = [("Offline-safe sync", "Every event arrives exactly once, SOS first after an outage."),
             ("SOS to incident", "One episode, one incident: owned, acknowledged, closed, audited."),
             ("Live monitoring", "One map of trips, devices and incidents, scoped by role."),
             ("Booking and rental", "From a booking to a provisioned device in the Guide's hands."),
             ("Return and billing", "Inspection, itemized fees, deposit first, sandbox payment."),
             ("Rules as configuration", "Every fee, threshold and window changeable live.")]
    names = [("Text 8", "Text 9"), ("Text 13", "Text 14"), ("Text 18", "Text 19"),
             ("Text 23", "Text 24"), ("Text 28", "Text 29"), ("Text 33", "Text 34")]
    for (t, d), (a, b) in zip(feats, names):
        set_text(shape(s5, a), t, size=13)
        set_text(shape(s5, b), d, size=12)

    # 6 Actors
    common(s6, 6, "Actors & Their Functions", "WHO USES THE SYSTEM", LONGNN)
    fill_table(s6, [
        ("Actor", "Key Functions"),
        ("Guest", "Browse packages · Register"),
        ("Customer", "Book · Pay · Sign agreement · View own history"),
        ("Operator (Staff)", "Confirm bookings · Allocate and provision devices · Check out and in · Inspect · Coordinate incidents"),
        ("Guide (Staff)", "Confirm handover · Start and end trip · Acknowledge incidents · Report from the field · Own trips only"),
        ("Admin (Staff)", "Users and roles · Pricing and parameters · Device catalogue · Audit and health"),
        ("Device, Gateway, Scheduler", "Emit SOS and positions · Buffer and deliver events · Expire holds, flag stale, escalate"),
    ])

    # Main Flows: steps (slide 7 pattern) + swimlane (slide 2 frame, boxes removed)
    order = [s1, s2, s3, s4, s5, s6]
    n = 7
    for key, title, who, steps, exc in MF:
        st = duplicate(prs, s7)
        common(st, n, title, "CORE SCENARIO", who)
        for (t, d), (a, b) in zip(steps, [("Text 8", "Text 9"), ("Text 13", "Text 14"), ("Text 18", "Text 19"),
                                           ("Text 23", "Text 24"), ("Text 27", "Text 28")]):
            set_text(shape(st, a), t, size=14)
            set_text(shape(st, b), d, size=12)
        ex = st.shapes.add_textbox(Mm(14), Mm(146), Mm(311), Mm(12))
        ex.text_frame.word_wrap = True
        ex.text_frame.text = "Exception: " + exc
        ex.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
        ex.text_frame.paragraphs[0].runs[0].font.bold = True
        order.append(st)
        n += 1
        im = duplicate(prs, s2)
        for nm in ("Shape 5", "Text 6", "Text 7", "Shape 8", "Text 9", "Text 10"):
            e = shape(im, nm)._element
            e.getparent().remove(e)
        common(im, n, title.split(":")[0] + " swimlane", "MAIN FLOW DIAGRAM", who)
        fig = OUT / f"assets/srs-fig{FIGS[key]}.png"
        from PIL import Image
        w, h = Image.open(fig).size
        box_w, box_h = 311.0, 134.0
        sc = min(box_w / w, box_h / h)
        pw, ph = w * sc, h * sc
        im.shapes.add_picture(str(fig), Mm(14 + (box_w - pw) / 2), Mm(31 + (box_h - ph) / 2), Mm(pw), Mm(ph))
        order.append(im)
        n += 1

    # Closing slides
    sc_ = duplicate(prs, s4)
    common(sc_, n, "Scope, stated now", "BOUNDARY", TAN)
    set_text(shape(sc_, "Text 6"), "In scope")
    set_text(shape(sc_, "Text 7"), ["Five Main Flows.", "58 use cases, 112 functional requirements (SRS §2, §3).",
                                    "Gateway with offline queue.", "Targeted firmware work: on-device queue, fleet channel key."], size=13)
    set_text(shape(sc_, "Text 9"), "Out of scope")
    set_text(shape(sc_, "Text 10"), ["Mesh-stack rearchitecture", "Native mobile apps", "Production payment",
                                     "Route recommendation", "Localization of the system", "Hardware certification", "Multi-tenancy"], size=13)
    order.append(sc_); n += 1

    me = duplicate(prs, s6)
    common(me, n, "How we will know it works", "MEASURED TARGETS", LONGLP)
    fill_table(me, [("Property", "Target"),
                    ("Gateway to cloud sync", "≤ 5 s when the uplink is available"),
                    ("Offline recovery", "≥ 99 % delivered after a 30 s to 30 min loss"),
                    ("Duplicate prevention", "0 duplicate incidents across a 10× replay"),
                    ("Priority ordering", "≥ 99 %: all P0 before any P2 or P3"),
                    ("Incident alert", "≤ 2 s from creation to connected clients"),
                    ("API latency", "≤ 300 ms at p95, 50 concurrent users")])
    order.append(me); n += 1

    ri = duplicate(prs, s6)
    common(ri, n, "Top risks", "RISK AND MITIGATION", HOANG)
    fill_table(ri, [("Risk (impact)", "Mitigation"),
                    ("SOS announced by one unacknowledged radio frame (Critical)", "Cadence-anomaly detection now; firmware marker on every beacon"),
                    ("Node queue drops the SOS first in an outage (Critical)", "On-device durable priority queue, Stage B"),
                    ("Map tiles misrepresenting national sovereignty (Critical)", "Goong Maps held in configuration; sovereignty screenshots to file"),
                    ("Few physical devices for radio testing (High)", "Simulation supplements, never presented as radio evidence")],
               widths=[140, 171])
    order.append(ri); n += 1

    pl = duplicate(prs, s6)
    common(pl, n, "Plan to Review 2", "GATES", TAN)
    fill_table(pl, [("Gate", "Assessed"),
                    ("Review 1, W4", "Problem, scope, requirement baseline"),
                    ("Review 2, W8", "Design baseline; MF-01 and MF-02 demonstrated"),
                    ("Faculty Council, W13", "Final product, main flows and 85 % of use cases"),
                    ("Submission and defence, W15", "Complete package")])
    order.append(pl); n += 1

    ask = duplicate(prs, s4)
    common(ask, n, "What we ask", "FEEDBACK", TAN)
    set_text(shape(ask, "Text 6"), "We ask the committee to")
    set_text(shape(ask, "Text 7"), ["1. Confirm the five Main Flows are the right scope for the term.",
                                    "2. Confirm the requirement baseline is clear and verifiable.",
                                    "3. Record feedback against an owner and a deadline; we close it with evidence at Review 2."], size=14)
    set_text(shape(ask, "Text 9"), "Thank you")
    set_text(shape(ask, "Text 10"), ["Questions."], size=20)
    order.append(ask)

    delete(prs, s7)
    delete(prs, s8)
    for i, sl in enumerate(order):
        move(prs, sl, i)
    out = OUT / "Review1_Slides_TrekLink.pptx"
    prs.save(str(out))
    return out


def render_outline() -> None:
    src = (Path(__file__).parent / "Review1_Slides_OUTLINE.src.md").read_text()
    out = re.sub(r"\{fig:([\w-]+)\}", lambda m: FIGS[m.group(1)], src)
    (OUT / "Review1_Slides_OUTLINE.md").write_text(out)


if __name__ == "__main__":
    render_outline()
    p = build(Path(sys.argv[1]))
    print(f"{p.name}: {len(Presentation(str(p)).slides)} slides")
